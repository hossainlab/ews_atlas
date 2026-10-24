"""Render preprint/manuscript.md to a submission-format PDF and DOCX.

    python preprint/build_manuscript.py

Figures are inserted where the text places them with ``[[FIGURE: name.png]]``,
each followed by its legend, taken verbatim from
``results/pub_figs/figure_captions.md`` so the legend in the manuscript and the
legend beside the figure can never drift apart. Supplementary legends are
appended as a separate section.

Fields that only the authors can supply - author list, lead contact, deposit DOI,
repository URL, contributions and declarations - are rendered as visible
``[TO BE COMPLETED]`` blocks rather than silently omitted, so that no one can
mistake an unfinished submission for a finished one.
"""
from __future__ import annotations

import os
import re
import sys
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (BaseDocTemplate, Frame, Image, PageBreak,
                                PageTemplate, Paragraph, Spacer)

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "preprint" / "manuscript.md"
CAPS = ROOT / "results" / "pub_figs" / "figure_captions.md"
FIGDIR = ROOT / "results" / "pub_figs"
OUT_PDF = ROOT / "preprint" / "EWSAtlas_manuscript.pdf"
OUT_SUPP = ROOT / "preprint" / "EWSAtlas_supplementary.pdf"
OUT_DOCX = ROOT / "preprint" / "EWSAtlas_manuscript.docx"

PLACEHOLDERS = {
    "AUTHORS": ("Author list, affiliations and corresponding author. "
                "Not supplied: author identity is not something this pipeline can "
                "or should infer."),
    "LEAD_CONTACT": "Name and e-mail of the lead contact.",
    "ZENODO_DOI": "Zenodo DOI. deposit/zenodo.json is prepared and needs creators "
                  "and a licence before upload.",
    "REPOSITORY_URL": "Public repository URL. The repository is ready but has not "
                      "been published anywhere.",
    "CONTRIBUTIONS": ("Acknowledgements, author contributions and declaration of "
                      "interests."),
}


def _clean(body: str) -> str:
    """Drop the markdown rules that separate caption sections.

    The caption file ends each figure block with a '---' rule; carried into the
    manuscript it printed as three dashes hanging off the end of the legend.
    """
    lines = [l for l in body.strip().splitlines() if l.strip().strip("-") != "" or l.strip() == ""]
    return "\n".join(lines).strip().rstrip("-").strip()


def legend_blocks() -> dict:
    """Map 'Figure 1' -> its legend text, from the shared caption file."""
    text = CAPS.read_text(encoding="utf-8")
    out, main = {}, re.split(r"\n## ", text)
    for block in main:
        m = re.match(r"(Figure [0-9S]+)[\.,]", block.strip())
        if m:
            body = block.split("\n", 1)[1] if "\n" in block else ""
            head = block.split("\n", 1)[0]
            out[m.group(1)] = (head, _clean(body))
    # supplementary legends live as bold paragraphs under one heading
    for m in re.finditer(r"\*\*(Figure S\d+)\.(.+?)\*\*(.*?)(?=\n\n\*\*Figure S|\Z)",
                         text, flags=re.S):
        out[m.group(1)] = (m.group(1) + "." + m.group(2), _clean(m.group(3)))
    return out


def styles():
    ss = getSampleStyleSheet()
    base = dict(fontName="Times-Roman", fontSize=11, leading=16.5,
                alignment=TA_JUSTIFY, spaceAfter=6)
    return {
        "title": ParagraphStyle("t", parent=ss["Normal"], fontName="Times-Bold",
                                fontSize=16, leading=20, spaceAfter=10,
                                alignment=TA_CENTER),
        "h1": ParagraphStyle("h1", parent=ss["Normal"], fontName="Times-Bold",
                             fontSize=13, leading=17, spaceBefore=14, spaceAfter=6),
        "h2": ParagraphStyle("h2", parent=ss["Normal"], fontName="Times-Bold",
                             fontSize=11.5, leading=15, spaceBefore=10, spaceAfter=4),
        "body": ParagraphStyle("b", parent=ss["Normal"], **base),
        "legend": ParagraphStyle("l", parent=ss["Normal"], fontName="Times-Roman",
                                 fontSize=9, leading=12, alignment=TA_JUSTIFY,
                                 spaceBefore=4, spaceAfter=12),
        "todo": ParagraphStyle("todo", parent=ss["Normal"], fontName="Times-Bold",
                               fontSize=10, leading=14, textColor=colors.HexColor("#8A1C1C"),
                               borderColor=colors.HexColor("#8A1C1C"), borderWidth=0.6,
                               borderPadding=5, spaceBefore=8, spaceAfter=10),
        "ref": ParagraphStyle("r", parent=ss["Normal"], fontName="Times-Roman",
                              fontSize=9.5, leading=13, leftIndent=12,
                              firstLineIndent=-12, spaceAfter=3),
    }


def inline_todo(s: str) -> str:
    """Render a placeholder that sits inside a sentence as visible, flagged text.

    Whole-line placeholders become their own red block; one embedded in a
    paragraph ("should be directed to [[LEAD_CONTACT]]") would otherwise reach
    the PDF as raw brackets, which is exactly the kind of thing that survives
    into a submitted file unnoticed.
    """
    return re.sub(r"\[\[(\w+)\]\]",
                  lambda m: "[TO BE COMPLETED: " + m.group(1).lower().replace("_", " ") + "]",
                  s)


def md_inline(s: str) -> str:
    s = inline_todo(s)
    s = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<i>\1</i>", s)
    s = re.sub(r"`([^`]+)`", r"<font face='Courier' size='9.5'>\1</font>", s)
    return s


def build():
    st = styles()
    legends = legend_blocks()
    text = SRC.read_text(encoding="utf-8")
    story, width = [], A4[0] - 50 * mm

    def para(buf, style="body"):
        if buf.strip():
            story.append(Paragraph(md_inline(buf.strip()), st[style]))

    buf = ""
    in_refs = False
    for raw in text.split("\n"):
        line = raw.rstrip()
        fig = re.match(r"\[\[FIGURE: (.+?)\]\]", line.strip())
        ph = re.match(r"\[\[([A-Z_]+)\]\]", line.strip())
        if fig or ph or line.startswith("#") or not line.strip():
            para(buf, "ref" if in_refs else "body"); buf = ""
        if fig:
            name = fig.group(1)
            path = FIGDIR / name
            if not path.exists():
                story.append(Paragraph(f"[MISSING FIGURE FILE: {name}]", st["todo"]))
                continue
            from PIL import Image as PILImage
            iw, ih = PILImage.open(path).size
            w = width
            h = w * ih / iw
            max_h = A4[1] - 70 * mm
            if h > max_h:
                h = max_h; w = h * iw / ih
            story.append(Spacer(1, 4))
            story.append(Image(str(path), width=w, height=h))
            key = re.match(r"Figure_(\d+|S\d+)_", name)
            if key:
                k = "Figure " + key.group(1)
                if k in legends:
                    head, body = legends[k]
                    story.append(Paragraph("<b>" + md_inline(head) + "</b> " +
                                           md_inline(body.replace("\n", " ")), st["legend"]))
            continue
        if ph:
            story.append(Paragraph("[TO BE COMPLETED] " + PLACEHOLDERS[ph.group(1)],
                                   st["todo"]))
            continue
        if line.startswith("### "):
            story.append(Paragraph(md_inline(line[4:]), st["h2"])); continue
        if line.startswith("## "):
            in_refs = line[3:].strip().lower().startswith("reference")
            story.append(Paragraph(md_inline(line[3:]), st["h1"])); continue
        if line.startswith("# "):
            story.append(Paragraph(md_inline(line[2:]), st["title"])); continue
        if not line.strip():
            para(buf, "ref" if in_refs else "body"); buf = ""; continue
        if in_refs and re.match(r"^\d+\. ", line.strip()):
            para(buf, "ref"); buf = line.strip()
        else:
            buf += (" " if buf else "") + line.strip()
    para(buf, "ref" if in_refs else "body")

    # supplementary legends, in order, after the main text
    sup = sorted((k for k in legends if k.startswith("Figure S")),
                 key=lambda s: int(s.split("S")[1]))
    if sup:
        story.append(PageBreak())
        story.append(Paragraph("Supplementary figure legends", st["h1"]))
        for k in sup:
            head, body = legends[k]
            story.append(Paragraph("<b>" + md_inline(head) + "</b> " +
                                   md_inline(body.replace("\n", " ")), st["legend"]))

    doc = BaseDocTemplate(str(OUT_PDF), pagesize=A4,
                          leftMargin=25 * mm, rightMargin=25 * mm,
                          topMargin=22 * mm, bottomMargin=22 * mm,
                          title="EWSAtlas manuscript")
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="f")

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Times-Roman", 8.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawCentredString(A4[0] / 2, 12 * mm, str(d.page))
        canvas.restoreState()

    doc.addPageTemplates([PageTemplate(id="p", frames=[frame], onPage=footer)])
    doc.build(story)
    print(f"wrote {OUT_PDF.name} ({OUT_PDF.stat().st_size/1e6:.2f} MB)")

    try:
        write_supplementary(legends)
        write_docx(text, legends)
    except Exception as exc:  # the PDF is the deliverable; DOCX is a convenience
        print(f"DOCX skipped: {type(exc).__name__}: {exc}")
    return 0


def write_supplementary(legends: dict) -> None:
    """Companion PDF: each supplementary figure followed by its legend.

    The main document lists the supplementary legends so a reader of the text
    alone knows what exists; the images travel in this file, which is how most
    journals and preprint servers want them.
    """
    from PIL import Image as PILImage

    st = styles()
    story = []
    width = A4[0] - 50 * mm
    story.append(Paragraph("Supplementary figures", st["h1"]))
    for path in sorted(FIGDIR.glob("Figure_S*.png"),
                       key=lambda p: int(re.match(r"Figure_S(\d+)_", p.name).group(1))):
        key = "Figure S" + re.match(r"Figure_S(\d+)_", path.name).group(1)
        iw, ih = PILImage.open(path).size
        w = width
        h = w * ih / iw
        max_h = A4[1] - 90 * mm
        if h > max_h:
            h = max_h
            w = h * iw / ih
        story.append(Spacer(1, 8))
        story.append(Image(str(path), width=w, height=h))
        if key in legends:
            head, body = legends[key]
            story.append(Paragraph("<b>" + md_inline(head) + "</b> " +
                                   md_inline(body.replace("\n", " ")), st["legend"]))
        else:
            raise KeyError(f"no legend for {key}")
        story.append(PageBreak())
    doc = BaseDocTemplate(str(OUT_SUPP), pagesize=A4, leftMargin=25 * mm,
                          rightMargin=25 * mm, topMargin=25 * mm,
                          bottomMargin=25 * mm, title="EWSAtlas supplementary figures")
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="f")
    doc.addPageTemplates([PageTemplate(id="p", frames=[frame])])
    doc.build(story)
    print(f"wrote {OUT_SUPP.name} ({OUT_SUPP.stat().st_size/1e6:.2f} MB)")


def write_docx(text: str, legends: dict) -> None:
    from docx import Document
    from docx.shared import Pt, Mm
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Mm(25)
    normal = d.styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(11)
    for raw in text.split("\n"):
        line = raw.strip()
        fig = re.match(r"\[\[FIGURE: (.+?)\]\]", line)
        ph = re.match(r"\[\[([A-Z_]+)\]\]", line)
        if fig:
            p = FIGDIR / fig.group(1)
            if p.exists():
                d.add_picture(str(p), width=Mm(160))
                key = re.match(r"Figure_(\d+|S\d+)_", fig.group(1))
                if key and "Figure " + key.group(1) in legends:
                    head, body = legends["Figure " + key.group(1)]
                    par = d.add_paragraph()
                    run = par.add_run(head + " " + body.replace("\n", " "))
                    run.font.size = Pt(9)
            continue
        if ph:
            par = d.add_paragraph()
            run = par.add_run("[TO BE COMPLETED] " + PLACEHOLDERS[ph.group(1)])
            run.bold = True
            continue
        if line.startswith("### "):
            d.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            d.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            d.add_heading(line[2:], level=1)
        elif line:
            d.add_paragraph(inline_todo(re.sub(r"[*`]", "", line)))
    tmp = OUT_DOCX.with_suffix(".tmp.docx")
    d.save(str(tmp))
    os.replace(tmp, OUT_DOCX)
    print(f"wrote {OUT_DOCX.name} ({OUT_DOCX.stat().st_size/1e6:.2f} MB)")


if __name__ == "__main__":
    raise SystemExit(build())
